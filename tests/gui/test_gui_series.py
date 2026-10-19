"""Reiter SERIE (Etappe 3e): Tabelle/Nummerierung, Spalten einfügen, Zeilen, Vorschau, Drucken."""
import os

import openpyxl
import pytest
from PySide6.QtCore import Qt

from labelwerk import storage
from labelwerk.document import Document, Numbering, Series, TextObject
from labelwerk.gui import main_window
from labelwerk.gui.main_window import MainWindow
from labelwerk.gui.tools import TOOLS
from labelwerk.printer import to_print_mode
from labelwerk.render import render_document
from labelwerk.series import fill
from labelwerk.settings import Settings

CSV = "Klemme;Bezeichnung\nX1;24V\nX2;0V\nX3;PE\n"


def _csv(tmp_path, text=CSV, name="klemmen.csv"):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def _window(qtbot, doc=None):
    win = MainWindow(doc=doc or Document("62x29", objects=[TextObject("{Klemme} {Bezeichnung}",
                                                                     size_pt=14)]),
                     settings=Settings(check_media_on_start=False))
    qtbot.addWidget(win)
    win.show()
    win.render_now()
    return win


def _with_table(qtbot, tmp_path, doc=None):
    win = _window(qtbot, doc)
    path = _csv(tmp_path)
    win.ask_series_file = lambda: str(path)
    win.panel("series").table_button.click()
    win.render_now()
    return win, path


class FakeWorker:
    """Druck-Worker, der nur mitschreibt (kein Thread, kein Drucker)."""
    jobs = []

    def __init__(self, image, label_id, model, cut_every=1, printer=None, options=None):
        FakeWorker.jobs.append((image, label_id, cut_every))
        from labelwerk.gui.print_worker import PrintSignals
        self.signals = PrintSignals()


@pytest.fixture
def printed(monkeypatch):
    FakeWorker.jobs = []
    monkeypatch.setattr(main_window, "PrintWorker", FakeWorker)
    monkeypatch.setattr(main_window.QThreadPool.globalInstance(), "start", lambda worker: None)
    return FakeWorker.jobs


# --- Reiter und Quelle -------------------------------------------------------------------------
def test_series_tab_follows_layout(qtbot):
    assert [tool.key for tool in TOOLS][-1] == "series"
    panel = _window(qtbot).panel("series")
    assert panel.source_label.text() == "Keine Serie"
    assert not panel.remove_button.isEnabled()


def test_choosing_a_table_links_it_and_lists_columns_and_rows(qtbot, tmp_path):
    win, path = _with_table(qtbot, tmp_path)
    panel = win.panel("series")
    assert win.doc.series == Series(source=str(path))
    assert panel.source_label.text() == "klemmen.csv"
    assert [panel.columns.item(i).text() for i in range(panel.columns.count())] == [
        "Klemme", "Bezeichnung"]
    assert panel.rows.rowCount() == 3 and panel.count_label.text() == "3 von 3 gewählt"
    assert panel.row_label.text() == "Zeile 1/3"


def test_unreadable_table_gives_a_message_and_changes_nothing(qtbot, tmp_path):
    win = _window(qtbot)
    bad = tmp_path / "bild.png"
    bad.write_bytes(b"x")
    win.ask_series_file = lambda: str(bad)
    win.panel("series").table_button.click()
    assert win.doc.series is None
    assert "Tabellenformat" in win.result.text()


def test_cancelled_dialog_changes_nothing(qtbot):
    win = _window(qtbot)
    win.ask_series_file = lambda: None
    win.panel("series").table_button.click()
    assert win.doc.series is None


def test_excel_sheets_can_be_chosen(qtbot, tmp_path):
    book = openpyxl.Workbook()
    book.active.title = "Erstes"
    book.active.append(["Klemme"])
    book.active.append(["A1"])
    second = book.create_sheet("Zweites")
    second.append(["Klemme"])
    second.append(["B1"])
    second.append(["B2"])
    path = tmp_path / "t.xlsx"
    book.save(path)
    win = _window(qtbot)
    win.ask_series_file = lambda: str(path)
    panel = win.panel("series")
    panel.table_button.click()
    assert panel.sheet.isVisibleTo(panel) and [panel.sheet.itemText(i) for i in range(2)] == [
        "Erstes", "Zweites"]
    panel.sheet.setCurrentIndex(1)
    assert win.doc.series.sheet == "Zweites" and panel.rows.rowCount() == 2


def test_numbering_creates_rows_and_fields_change_it(qtbot):
    win = _window(qtbot, Document("62x29", objects=[TextObject("X{Nummer}", size_pt=14)]))
    panel = win.panel("series")
    panel.numbering_button.click()
    assert win.doc.series == Series(numbering=Numbering())
    assert panel.rows.rowCount() == 10 and panel.numbering_box.isVisibleTo(panel)
    panel.count.setValue(3)
    panel.digits.setValue(2)
    assert win.doc.series.numbering == Numbering(count=3, digits=2)
    win.render_now()
    assert win.image.tobytes() == render_document(
        fill(win.doc, ["Nummer"], ["01"])).tobytes()
    win.shortcut_actions["edit.undo"].trigger()
    assert win.doc.series.numbering == Numbering(count=3) and panel.rows.rowCount() == 3


def test_step_zero_is_not_taken(qtbot):
    win = _window(qtbot)
    panel = win.panel("series")
    panel.numbering_button.click()
    panel.step.setValue(0)
    assert win.doc.series.numbering.step == 1
    assert "Schritt" in panel.error_label.text()


def test_remove_takes_the_series_off(qtbot, tmp_path):
    win, _ = _with_table(qtbot, tmp_path)
    win.panel("series").remove_button.click()
    assert win.doc.series is None and win.panel("series").rows.rowCount() == 0


def test_cut_mode_is_saved_in_the_series(qtbot, tmp_path):
    win, _ = _with_table(qtbot, tmp_path)
    panel = win.panel("series")
    panel.cut.setCurrentIndex(panel.cut.findData("end"))
    assert win.doc.series.cut == "end"
    win.shortcut_actions["edit.undo"].trigger()
    assert win.doc.series.cut == "each" and panel.cut.currentData() == "each"


# --- Vorschau ----------------------------------------------------------------------------------
def test_canvas_shows_the_preview_row(qtbot, tmp_path):
    win, _ = _with_table(qtbot, tmp_path)
    columns = ["Klemme", "Bezeichnung"]
    assert win.image.tobytes() == render_document(fill(win.doc, columns, ["X1", "24V"])).tobytes()
    assert win.doc.objects[0].text == "{Klemme} {Bezeichnung}"        # Etikett unverändert
    panel = win.panel("series")
    panel.next_button.click()
    win.render_now()
    assert panel.row_label.text() == "Zeile 2/3"
    assert win.image.tobytes() == render_document(fill(win.doc, columns, ["X2", "0V"])).tobytes()
    panel.rows.setCurrentCell(2, 0)
    win.render_now()
    assert panel.row_label.text() == "Zeile 3/3"
    assert not panel.next_button.isEnabled()


def test_writing_mode_shows_the_placeholders(qtbot, tmp_path):
    win, _ = _with_table(qtbot, tmp_path, Document("62x29", objects=[
        TextObject("{Klemme} {Bezeichnung}", size_pt=8)]))
    win.edit_text(0)
    win.render_now()
    assert win.rendered[0].layout.texts == ["{Klemme} {Bezeichnung}"]
    win.finish_editing()
    win.render_now()
    assert win.rendered[0].layout.texts == ["X1 24V"]


def test_unknown_placeholder_is_pointed_out(qtbot, tmp_path):
    win, _ = _with_table(qtbot, tmp_path, Document("62x29", objects=[TextObject("{Foo}")]))
    assert "Unbekannte Spalte: Foo" in win.panel("series").error_label.text()


def test_missing_table_after_opening_is_shown(qtbot, tmp_path):
    path = _csv(tmp_path)
    doc = Document("62x29", objects=[TextObject("{Klemme}")], series=Series(source=str(path)))
    saved = storage.save(doc, tmp_path / "s.labelwerk")
    os.remove(path)
    win = _window(qtbot)
    win.files.open(saved, ask=False)
    win.render_now()
    panel = win.panel("series")
    assert win.doc.series == Series(source=str(path))
    assert "nicht gefunden" in panel.error_label.text()
    assert "nicht gefunden" in win.result.text()


def _touch(path, text):
    path.write_text(text, encoding="utf-8")
    stamp = os.stat(path).st_mtime + 5
    os.utime(path, (stamp, stamp))


def test_changed_table_file_is_read_again_when_checked(qtbot, tmp_path):
    win, path = _with_table(qtbot, tmp_path)
    panel = win.panel("series")
    assert not panel.check_file()
    _touch(path, "Klemme;Bezeichnung\nY1;5V\n")
    panel.refresh()                                     # z. B. Rückgängig: liest nicht neu
    assert panel.rows.rowCount() == 3
    assert panel.check_file() and panel.rows.rowCount() == 1


# --- Spalten einfügen --------------------------------------------------------------------------
def test_column_is_appended_to_the_selected_text(qtbot, tmp_path):
    win, _ = _with_table(qtbot, tmp_path, Document("62x29", objects=[TextObject("Nr. ", size_pt=14)]))
    win.canvas.set_selection([0])
    win.panel("series").insert_column.emit("Klemme")
    assert win.doc.objects[0].text == "Nr. {Klemme}"
    win.shortcut_actions["edit.undo"].trigger()
    assert win.doc.objects[0].text == "Nr. "


def test_column_without_selection_becomes_a_new_text(qtbot, tmp_path):
    win, _ = _with_table(qtbot, tmp_path, Document("62x29", objects=[]))
    win.panel("series").insert_column.emit("Bezeichnung")
    assert [o.text for o in win.doc.objects] == ["{Bezeichnung}"]
    assert win.canvas.selection() == [0]


def test_double_click_on_a_column_inserts_it_once(qtbot, tmp_path):
    from PySide6.QtTest import QTest
    win, _ = _with_table(qtbot, tmp_path, Document("62x29", objects=[]))
    win.tool_actions["series"].trigger()
    panel = win.panel("series")
    viewport = panel.columns.viewport()
    center = panel.columns.visualItemRect(panel.columns.item(0)).center()
    left = Qt.MouseButton.LeftButton
    QTest.mouseClick(viewport, left, pos=center)          # echte Folge eines Doppelklicks
    QTest.mousePress(viewport, left, pos=center)
    QTest.mouseDClick(viewport, left, pos=center)
    QTest.mouseRelease(viewport, left, pos=center)
    assert [o.text for o in win.doc.objects] == ["{Klemme}"]


# --- Drucken ----------------------------------------------------------------------------------
def test_printing_a_series_prints_the_checked_rows_in_one_job(qtbot, tmp_path, printed):
    win, _ = _with_table(qtbot, tmp_path)
    panel = win.panel("series")
    panel.rows.item(1, 0).setCheckState(Qt.CheckState.Unchecked)
    assert panel.count_label.text() == "2 von 3 gewählt"
    asked = []
    win.ask_print_count = lambda n: asked.append(n) or True
    win.print_label()
    assert asked == [2]
    (images, label, cut_every), = printed
    columns = ["Klemme", "Bezeichnung"]
    assert [i.tobytes() for i in images] == [
        to_print_mode(render_document(fill(win.doc, columns, row))).tobytes()
        for row in (["X1", "24V"], ["X3", "PE"])]
    assert (label, cut_every) == ("62x29", 1)


def test_cut_only_at_the_end_is_passed_on(qtbot, tmp_path, printed):
    win, _ = _with_table(qtbot, tmp_path)
    panel = win.panel("series")
    panel.cut.setCurrentIndex(panel.cut.findData("end"))
    win.ask_print_count = lambda n: True
    win.print_label()
    assert printed[0][2] == 3


def test_no_checked_row_prints_nothing(qtbot, tmp_path, printed):
    win, _ = _with_table(qtbot, tmp_path)
    win.panel("series").none_button.click()
    win.print_label()
    assert printed == [] and "Keine Zeile" in win.result.text()


def test_declined_count_prints_nothing(qtbot, tmp_path, printed):
    win, _ = _with_table(qtbot, tmp_path)
    win.ask_print_count = lambda n: False
    win.print_label()
    assert printed == []


def test_row_that_does_not_fit_stops_printing_with_its_number(qtbot, tmp_path, printed):
    win = _window(qtbot, Document("62x29", objects=[TextObject("{Bezeichnung}", size_pt=40)]))
    path = _csv(tmp_path, "Klemme;Bezeichnung\nX1;kurz\nX2;" + "sehr lang " * 30 + "\n")
    win.ask_series_file = lambda: str(path)
    win.panel("series").table_button.click()
    win.ask_print_count = lambda n: True
    win.print_label()
    assert printed == [] and "Zeile 2" in win.result.text()


def test_missing_table_prints_nothing(qtbot, tmp_path, printed):
    win, path = _with_table(qtbot, tmp_path)
    os.remove(path)
    win.panel("series").refresh()
    win.print_label()
    assert printed == [] and "nicht gefunden" in win.result.text()


def test_favorite_with_series_prints_all_rows(qtbot, tmp_path, printed):
    path = _csv(tmp_path)
    doc = Document("62x29", objects=[TextObject("{Klemme}")], series=Series(source=str(path)))
    win = _window(qtbot)
    win.ask_print_count = lambda n: n == 3
    win.print_document(doc)
    assert len(printed) == 1 and len(printed[0][0]) == 3


# --- Ansicht --------------------------------------------------------------------------------------
def test_preview_row_is_selected_in_the_table(qtbot, tmp_path):
    win, _ = _with_table(qtbot, tmp_path)
    panel = win.panel("series")
    panel.next_button.click()
    assert panel.rows.currentRow() == 1


@pytest.mark.parametrize("theme", ["rawblock", "classic"])
def test_series_panel_fits_its_width(qtbot, tmp_path, theme):
    from PySide6.QtWidgets import QApplication
    from labelwerk.gui.app import apply_theme
    apply_theme(QApplication.instance(), theme)
    try:
        win, _ = _with_table(qtbot, tmp_path)
        win.resize(1200, 980)
        win.tool_actions["series"].trigger()
        QApplication.processEvents()
        panel = win.panel("series")
        for button in (panel.table_button, panel.numbering_button, panel.remove_button,
                       panel.all_button, panel.none_button):
            assert button.width() >= button.sizeHint().width(), button.text()
        assert panel.count_label.width() >= panel.count_label.sizeHint().width()
        assert not panel.rows.horizontalScrollBar().isVisible()
    finally:
        apply_theme(QApplication.instance(), "classic")


# --- Nachbesserungen nach dem Review --------------------------------------------------------------
def _flush(qtbot, win):
    """Warten, bis ein geplantes Rendern gelaufen ist (ohne render_now von Hand)."""
    qtbot.waitUntil(lambda: not win._timer.isActive(), timeout=2000)


def test_writing_mode_shows_placeholders_and_caret_without_manual_render(qtbot, tmp_path):
    doc = Document("62x29", objects=[TextObject("{Klemme} {Bezeichnung}", size_pt=8, x_mm=5,
                                                y_mm=5)])
    win, _ = _with_table(qtbot, tmp_path, doc)
    win.edit_text(0)
    assert win.rendered[0].layout.texts == ["{Klemme} {Bezeichnung}"]
    assert win.canvas._caret is not None
    win.finish_editing()
    _flush(qtbot, win)
    assert win.rendered[0].layout.texts == ["X1 24V"]


@pytest.mark.parametrize("align", ["center", "right"])
def test_position_stays_over_a_writing_cycle(qtbot, tmp_path, align):
    from labelwerk.document import Paragraph, Span
    obj = TextObject([Paragraph([Span("{Bezeichnung}")], align)], size_pt=8, x_mm=10, y_mm=5)
    win, path = _with_table(qtbot, tmp_path, Document("62x29", objects=[obj]))
    _touch(path, "Klemme;Bezeichnung\nX1;Ein sehr langer Wert hier\n")
    win.panel("series").check_file()
    win.render_now()
    before = (win.doc.objects[0].x_mm, win.doc.objects[0].y_mm)
    for _ in range(3):
        win.edit_text(0)
        win.panel("text").edit.textCursor().insertText("x")
        _flush(qtbot, win)
        win.panel("text").edit.textCursor().deletePreviousChar()
        _flush(qtbot, win)
        win.finish_editing()
        _flush(qtbot, win)
    assert (win.doc.objects[0].x_mm, win.doc.objects[0].y_mm) == before


def test_too_long_placeholder_while_writing_is_explained(qtbot, tmp_path):
    win, _ = _with_table(qtbot, tmp_path, Document("62x29", objects=[
        TextObject("{Bezeichnung}", size_pt=40)]))
    win.edit_text(0)
    assert "Platzhalter" in win.result.text()


def test_render_error_of_the_preview_row_names_it(qtbot, tmp_path):
    win = _window(qtbot, Document("62x29", objects=[TextObject("{Bezeichnung}", size_pt=40)]))
    path = _csv(tmp_path, "Klemme;Bezeichnung\nX1;kurz\nX2;" + "sehr lang " * 30 + "\n")
    win.ask_series_file = lambda: str(path)
    win.panel("series").table_button.click()
    win.panel("series").next_button.click()
    win.render_now()
    assert win.result.text().startswith("Zeile 2:")


def test_numbering_error_clears_when_fixed(qtbot):
    win = _window(qtbot)
    panel = win.panel("series")
    panel.numbering_button.click()
    panel.step.setValue(0)
    assert "Schritt" in panel.error_label.text()
    panel.step.setValue(1)
    assert not panel.error_label.isVisibleTo(panel)
    panel.start.setValue(10**9)                        # letzte Zahl zu groß: nicht übernommen
    assert win.doc.series.numbering.start == 1 and "zu groß" in panel.error_label.text()
    panel.refresh()
    assert not panel.error_label.isVisibleTo(panel)


def test_series_tab_stays_when_objects_are_clicked(qtbot, tmp_path):
    win, _ = _with_table(qtbot, tmp_path, Document("62x29", objects=[]))
    win.tool_actions["series"].trigger()
    panel = win.panel("series")
    panel.insert_column.emit("Klemme")
    panel.insert_column.emit("Bezeichnung")
    assert win.panels.currentWidget() is panel
    assert [o.text for o in win.doc.objects] == ["{Klemme} {Bezeichnung}"]


def test_rows_that_stick_out_are_asked_about(qtbot, tmp_path, printed):
    doc = Document("62x29", objects=[TextObject("{Bezeichnung}", size_pt=14, x_mm=30, y_mm=5)])
    win = _window(qtbot, doc)
    path = _csv(tmp_path, "Klemme;Bezeichnung\nX1;kurz\nX2;" + "lang " * 12 + "\nX3;ok\n")
    win.ask_series_file = lambda: str(path)
    win.panel("series").table_button.click()
    win.render_now()
    win.ask_print_count = lambda n: True
    asked = []
    win.confirm_outside_rows = lambda numbers: asked.append(numbers) or False
    win.print_label()
    assert asked == [[2]] and printed == []


def test_count_is_asked_before_rendering_and_media_check_first(qtbot, tmp_path, printed,
                                                               monkeypatch):
    win, _ = _with_table(qtbot, tmp_path)
    rendered = []
    real = main_window.Renderer.render
    monkeypatch.setattr(main_window.Renderer, "render",
                        lambda self, doc, model="QL-600", **k: rendered.append(1)
                        or real(self, doc, model, **k))
    win.render_now()
    rendered.clear()
    win.ask_print_count = lambda n: False
    win.print_label()
    assert rendered == []
    win.is_checking_media = lambda: True
    win.ask_print_count = lambda n: pytest.fail("gefragt, obwohl die Rolle geprüft wird")
    win.print_label()
    assert "Rolle wird gerade geprüft" in win.result.text()


def test_changed_table_stops_printing_with_a_hint(qtbot, tmp_path, printed):
    win, path = _with_table(qtbot, tmp_path)
    _touch(path, CSV + "X4;N\n")
    win.ask_print_count = lambda n: pytest.fail("gefragt")
    win.print_label()
    assert printed == [] and "neu eingelesen" in win.result.text()
    assert win.panel("series").rows.rowCount() == 4


def test_opening_resets_the_checks(qtbot, tmp_path):
    path = _csv(tmp_path)
    doc = Document("62x29", objects=[TextObject("{Klemme}")], series=Series(source=str(path)))
    saved = storage.save(doc, tmp_path / "s.labelwerk")
    win = _window(qtbot)
    win.files.open(saved, ask=False)
    panel = win.panel("series")
    panel.rows.item(0, 0).setCheckState(Qt.CheckState.Unchecked)
    panel.next_button.click()
    win.files.open(saved, ask=False)
    assert panel.count_label.text() == "3 von 3 gewählt" and panel.row_label.text() == "Zeile 1/3"


def test_missing_sheet_still_offers_the_others(qtbot, tmp_path):
    book = openpyxl.Workbook()
    book.active.title = "Daten"
    book.active.append(["A"])
    book.active.append([1])
    path = tmp_path / "t.xlsx"
    book.save(path)
    win = _window(qtbot, Document("62x29", objects=[TextObject("{A}")],
                                  series=Series(source=str(path), sheet="Weg")))
    panel = win.panel("series")
    assert "Blatt" in panel.error_label.text()
    assert panel.sheet.isVisibleTo(panel) and panel.sheet.itemText(0) == "Daten"


def test_table_without_rows_says_so(qtbot, tmp_path, printed):
    win = _window(qtbot)
    path = _csv(tmp_path, "Klemme;Bezeichnung\n")
    win.ask_series_file = lambda: str(path)
    win.panel("series").table_button.click()
    win.print_label()
    assert printed == [] and "keine Zeilen" in win.result.text()


def test_table_is_read_once_when_chosen(qtbot, tmp_path, monkeypatch):
    from labelwerk import series as core
    win = _window(qtbot)
    path = _csv(tmp_path)
    calls = []
    real = core.load_table
    monkeypatch.setattr(core, "load_table", lambda *a, **k: calls.append(1) or real(*a, **k))
    win.ask_series_file = lambda: str(path)
    win.panel("series").table_button.click()
    assert len(calls) == 1


def test_file_dialog_starts_at_the_current_table(qtbot, tmp_path, monkeypatch):
    win, path = _with_table(qtbot, tmp_path)
    seen = []
    monkeypatch.setattr(main_window.QFileDialog, "getOpenFileName",
                        lambda parent, title, folder, filters: seen.append(folder) or ("", ""))
    del win.ask_series_file
    win.choose_series_table()
    assert seen == [str(path.parent)]


# --- Schlussreview --------------------------------------------------------------------------------
def _hidden_first(tmp_path):
    book = openpyxl.Workbook()
    hidden = book.active
    hidden.title = "Versteckt"
    hidden.append(["Klemme"])
    hidden.append(["HIDDEN"])
    hidden.sheet_state = "hidden"
    shown = book.create_sheet("Sichtbar")
    shown.append(["Klemme"])
    shown.append(["S1"])
    path = tmp_path / "t.xlsx"
    book.save(path)
    return path


def test_choosing_excel_reads_once_and_keeps_the_visible_sheet(qtbot, tmp_path, monkeypatch):
    path = _hidden_first(tmp_path)
    opened = []
    real = openpyxl.load_workbook
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: opened.append(1) or real(*a, **k))
    win = _window(qtbot, Document("62x29", objects=[TextObject("{Klemme}")]))
    win.ask_series_file = lambda: str(path)
    panel = win.panel("series")
    panel.table_button.click()
    assert len(opened) == 1
    assert win.doc.series.sheet is None                  # erst eine eigene Wahl legt das Blatt fest
    assert panel.sheet.currentText() == "Sichtbar"
    assert panel.preview()[1] == ["S1"]


def test_opening_reads_the_table_once(qtbot, tmp_path, monkeypatch):
    from labelwerk import series as core
    path = _csv(tmp_path)
    doc = Document("62x29", objects=[TextObject("{Klemme}")], series=Series(source=str(path)))
    saved = storage.save(doc, tmp_path / "s.labelwerk")
    win = _window(qtbot)
    calls = []
    real = core.load_table
    monkeypatch.setattr(core, "load_table", lambda *a, **k: calls.append(1) or real(*a, **k))
    win.files.open(saved, ask=False)
    assert len(calls) == 1


def test_centred_text_stays_centred_for_every_row(qtbot, tmp_path):
    from labelwerk import layout
    from labelwerk.document import Paragraph, Span
    obj = TextObject([Paragraph([Span("{Klemme}")], "center")], size_pt=12, x_mm=25, y_mm=5)
    win = _window(qtbot, Document("62x29", objects=[obj]))
    path = _csv(tmp_path, "Klemme;B\nX1;1\nX1 -K3 Klemme 24V;2\n")
    win.ask_series_file = lambda: str(path)
    win.panel("series").table_button.click()
    win.render_now()
    centres = [layout.object_rect(win.rendered[0]).cx]
    win.panel("series").next_button.click()
    win.render_now()
    centres.append(layout.object_rect(win.rendered[0]).cx)
    assert centres[1] == pytest.approx(centres[0], abs=0.2)
