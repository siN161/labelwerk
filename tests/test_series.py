"""Series (stage 3e): tables from CSV/Excel, numbering, placeholders."""
import datetime
import os

import openpyxl
import pytest

from labelwerk import series, storage
from labelwerk.document import Document, ImageObject, Numbering, Paragraph, Series, Span, TextObject
from labelwerk.series import (SeriesError, Table, fill, load_table, missing, numbering_table,
                              placeholders, sheet_names, table_for)


def _csv(tmp_path, text, encoding="utf-8", name="t.csv"):
    path = tmp_path / name
    path.write_bytes(text.encode(encoding))
    return path


def _xlsx(tmp_path, sheets, name="t.xlsx"):
    """Excel file with sheets ``{name: rows}`` (first row = column names)."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    path = tmp_path / name
    book.save(path)
    return path


# --- CSV -------------------------------------------------------------------------------------
def test_csv_with_semicolons(tmp_path):
    table = load_table(_csv(tmp_path, "Klemme;Bezeichnung\nX1;24V\nX2;0V\n"))
    assert table == Table(["Klemme", "Bezeichnung"], [["X1", "24V"], ["X2", "0V"]])


@pytest.mark.parametrize("sep", [",", "\t", "|"])
def test_csv_separator_is_detected(tmp_path, sep):
    table = load_table(_csv(tmp_path, f"A{sep}B\n1{sep}2\n"))
    assert table == Table(["A", "B"], [["1", "2"]])


def test_semicolon_wins_over_decimal_commas(tmp_path):
    table = load_table(_csv(tmp_path, "Wert;Einheit\n1,5;mm\n2,25;mm\n"))
    assert table.rows == [["1,5", "mm"], ["2,25", "mm"]]


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1252"])
def test_csv_encodings(tmp_path, encoding):
    table = load_table(_csv(tmp_path, "Straße;Größe\nMaß;ä\n", encoding))
    assert table == Table(["Straße", "Größe"], [["Maß", "ä"]])


def test_csv_quotes_and_line_breaks_in_cells(tmp_path):
    table = load_table(_csv(tmp_path, 'A;B\n"x;y";"Zeile\neins"\n'))
    assert table.rows == [["x;y", "Zeile eins"]]


def test_empty_and_duplicate_column_names(tmp_path):
    table = load_table(_csv(tmp_path, ";Name;Name\n1;2;3\n"))
    assert table.columns == ["Spalte 1", "Name", "Name (2)"]


def test_short_rows_are_padded_long_ones_cut_empty_ones_dropped(tmp_path):
    table = load_table(_csv(tmp_path, "A;B\n1\n\n;\n2;3;4\n\n"))
    assert table.rows == [["1", ""], ["2", "3"]]


def test_csv_limits(tmp_path, monkeypatch):
    monkeypatch.setattr(series, "MAX_ROWS", 2)
    with pytest.raises(SeriesError, match="zu viele Zeilen"):
        load_table(_csv(tmp_path, "A\n1\n2\n3\n"))
    monkeypatch.setattr(series, "MAX_COLUMNS", 2)
    with pytest.raises(SeriesError, match="zu viele Spalten"):
        load_table(_csv(tmp_path, "A;B;C\n1;2;3\n"))
    monkeypatch.setattr(series, "MAX_BYTES", 5)
    with pytest.raises(SeriesError, match="zu groß"):
        load_table(_csv(tmp_path, "A;B\n1;2\n"))


def test_missing_empty_and_unknown_files(tmp_path):
    with pytest.raises(SeriesError, match="nicht gefunden"):
        load_table(tmp_path / "fehlt.csv")
    with pytest.raises(SeriesError, match="leer"):
        load_table(_csv(tmp_path, ""))
    with pytest.raises(SeriesError, match="Tabellenformat"):
        load_table(_csv(tmp_path, "A;B\n", name="t.pdf"))
    with pytest.raises(SeriesError, match="nicht lesbar"):
        load_table(_csv(tmp_path, "A;B\n1\x002\n"))
    folder = tmp_path / "ordner.csv"
    folder.mkdir()
    with pytest.raises(SeriesError, match="ordner.csv"):
        load_table(folder)


# --- Excel -----------------------------------------------------------------------------------
def test_xlsx_values_are_shown_like_in_excel(tmp_path):
    path = _xlsx(tmp_path, {"Daten": [
        ["Nr", "Wert", "Datum", "Leer"],
        [1, 1.5, datetime.datetime(2026, 10, 18), None],
        [2.0, "Text", datetime.datetime(2026, 10, 18, 14, 30), "x"],
        [True, 0.1 + 0.2, datetime.date(2026, 1, 2), "a\nb"],
    ]})
    assert load_table(path) == Table(["Nr", "Wert", "Datum", "Leer"], [
        ["1", "1,5", "18.10.2026", ""],
        ["2", "Text", "18.10.2026 14:30", "x"],
        ["WAHR", "0,3", "02.01.2026", "a b"],
    ])


def test_xlsx_sheets(tmp_path):
    path = _xlsx(tmp_path, {"Erstes": [["A"], [1]], "Zweites": [["B"], [2]]})
    assert sheet_names(path) == ["Erstes", "Zweites"]
    assert load_table(path) == Table(["A"], [["1"]])
    assert load_table(path, "Zweites") == Table(["B"], [["2"]])
    with pytest.raises(SeriesError, match="Blatt „Drittes“"):
        load_table(path, "Drittes")
    assert sheet_names(_csv(tmp_path, "A\n1\n")) == []


def test_broken_xlsx_gives_a_message(tmp_path):
    path = tmp_path / "kaputt.xlsx"
    path.write_bytes(b"PK\x03\x04 kein richtiges Archiv")
    with pytest.raises(SeriesError, match="kaputt.xlsx.*nicht lesbar"):
        load_table(path)


def test_xlsx_unpacked_size_is_limited(tmp_path, monkeypatch):
    path = _xlsx(tmp_path, {"Daten": [["A"], ["x" * 1000]]})
    monkeypatch.setattr(series, "MAX_UNPACKED", 1000)
    with pytest.raises(SeriesError, match="zu groß"):
        load_table(path)


def test_excel_xml_is_read_safely():
    assert openpyxl.DEFUSEDXML            # defusedxml protects against XML attacks in foreign files


# --- Numbering -------------------------------------------------------------------------------
def test_numbering_counts_up():
    assert numbering_table(Numbering(start=1, step=1, count=3)) == Table(
        ["Nummer"], [["1"], ["2"], ["3"]])


def test_numbering_with_step_and_leading_zeros():
    assert numbering_table(Numbering(start=8, step=2, count=3, digits=3)).rows == [
        ["008"], ["010"], ["012"]]
    assert numbering_table(Numbering(start=2, step=-1, count=4, digits=2)).rows == [
        ["02"], ["01"], ["00"], ["-01"]]


@pytest.mark.parametrize("numbering, message", [
    (Numbering(count=0), "Anzahl"), (Numbering(count=10_001), "Anzahl"),
    (Numbering(step=0), "Schritt"), (Numbering(start=10**10), "Start"),
    (Numbering(digits=11), "Stellen"),
])
def test_invalid_numbering(numbering, message):
    with pytest.raises(SeriesError, match=message):
        numbering_table(numbering)


def test_table_for_series(tmp_path):
    path = _csv(tmp_path, "A\n1\n")
    assert table_for(Series(source=str(path))) == Table(["A"], [["1"]])
    assert table_for(Series(source="t.csv"), base=tmp_path) == Table(["A"], [["1"]])
    assert table_for(Series(numbering=Numbering(count=2))).rows == [["1"], ["2"]]


@pytest.mark.parametrize("value, message", [
    (Series(), "Tabelle oder Nummerierung"),
    (Series(source="a.csv", numbering=Numbering()), "Tabelle oder Nummerierung"),
    (Series(numbering=Numbering(), cut="nie"), "Schnitt"),
])
def test_invalid_series(value, message):
    with pytest.raises(SeriesError, match=message):
        series.check(value)


# --- Placeholders --------------------------------------------------------------------------------
def _doc(*objects):
    return Document("62x29", objects=list(objects))


def test_placeholders_in_order_once():
    doc = _doc(TextObject("{Klemme}\n{Bezeichnung} V {Klemme}"), TextObject("{Ort}"))
    assert placeholders(doc) == ["Klemme", "Bezeichnung", "Ort"]


def test_fill_replaces_known_columns_and_keeps_the_original():
    doc = _doc(TextObject("{Klemme}\n{Bezeichnung} V"))
    filled = fill(doc, ["Klemme", "Bezeichnung"], ["X1", "24"])
    assert filled.objects[0].text == "X1\n24 V"
    assert doc.objects[0].text == "{Klemme}\n{Bezeichnung} V"


def test_fill_keeps_the_formatting_around_and_of_the_placeholder():
    obj = TextObject([Paragraph([Span("Nr. "), Span("{Klemme}", italic=True), Span(" ok")])])
    spans = fill(_doc(obj), ["Klemme"], ["X1"]).objects[0].paragraphs[0].spans
    assert [(s.text, s.italic) for s in spans] == [("Nr. ", False), ("X1", True), (" ok", False)]


def test_fill_across_spans_takes_the_formatting_where_it_starts():
    obj = TextObject([Paragraph([Span("A {Kl", bold=False), Span("emme} B", underline=True)],
                                "left")])
    para = fill(_doc(obj), ["Klemme"], ["X1"]).objects[0].paragraphs[0]
    assert para.align == "left"
    assert [(s.text, s.bold, s.underline) for s in para.spans] == [
        ("A X1", False, False), (" B", True, True)]


def test_fill_whole_paragraph_keeps_one_span():
    obj = TextObject([Paragraph([Span("{Leer}", size_pt=20)])])
    para = fill(_doc(obj), ["Leer"], [""]).objects[0].paragraphs[0]
    assert [(s.text, s.size_pt) for s in para.spans] == [("", 20)]


def test_unknown_placeholders_stay_and_are_reported():
    doc = _doc(TextObject("{Klemme} {Foo}"))
    assert fill(doc, ["Klemme"], ["X1"]).objects[0].text == "X1 {Foo}"
    assert missing(doc, ["Klemme"]) == ["Foo"]


def test_values_are_inserted_once_and_without_line_breaks():
    doc = _doc(TextObject("{A}|{B}"))
    assert fill(doc, ["A", "B"], ["{B}", "x\ny"]).objects[0].text == "{B}|x y"


def test_fill_keeps_one_object_and_images():
    image = ImageObject.from_image(__import__("PIL.Image").Image.new("RGB", (10, 10)), 300, "b.png")
    doc = _doc(TextObject("{A}"), TextObject("{A}"), image)
    filled = fill(doc, ["A"], ["1"], keep=1)
    assert [o.text for o in filled.objects[:2]] == ["1", "{A}"]
    assert filled.objects[2] == image


# --- Saving ----------------------------------------------------------------------------------
@pytest.mark.parametrize("value", [
    Series(numbering=Numbering(start=5, step=5, count=20, digits=3), cut="end"),
    Series(source="/tmp/klemmen.xlsx", sheet="Blatt 2"),
])
def test_series_roundtrip(tmp_path, value):
    doc = Document("62", objects=[TextObject("X{Nummer}")], series=value)
    path = storage.save(doc, tmp_path / "s.labelwerk")
    assert storage.load(path) == doc


def test_document_without_series_saves_null(tmp_path):
    import json
    path = storage.save(Document("62"), tmp_path / "s.labelwerk")
    assert json.loads(path.read_text())["document"]["series"] is None
    assert storage.load(path).series is None


@pytest.mark.parametrize("value, message", [
    ("viele", "Serie"),
    ({"cut": "each"}, "Tabelle oder Nummerierung"),
    ({"source": "a.csv", "numbering": {"count": 2}}, "Tabelle oder Nummerierung"),
    ({"source": 3}, "source"),
    ({"source": "a.csv", "cut": "nie"}, "cut"),
    ({"numbering": {"count": 0}}, "count"),
    ({"numbering": {"step": 0}}, "step"),
    ({"numbering": {"start": 1.5}}, "start"),
    ({"numbering": {"digits": 99}}, "digits"),
    ({"numbering": "drei"}, "Nummerierung"),
])
def test_invalid_series_in_files(tmp_path, value, message):
    import json
    data = storage.to_dict(Document("62"))
    data["document"]["series"] = value
    path = tmp_path / "s.labelwerk"
    path.write_text(json.dumps(data))
    with pytest.raises(storage.StorageError, match=message):
        storage.load(path)


# --- Improvements after the review ------------------------------------------------------------
import io  # noqa: E402
import zipfile  # noqa: E402


def test_excel_rows_after_a_gap_are_not_lost(tmp_path, monkeypatch):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["A1"], sheet["A2"], sheet["A40"] = "Klemme", "X1", "spät"
    path = tmp_path / "luecke.xlsx"
    book.save(path)
    assert load_table(path).rows == [["X1"], ["spät"]]
    monkeypatch.setattr(series, "MAX_ROWS", 1)
    with pytest.raises(SeriesError, match="zu viele Zeilen"):
        load_table(path)


def test_exactly_max_rows_is_fine(tmp_path, monkeypatch):
    monkeypatch.setattr(series, "MAX_ROWS", 2)
    assert len(load_table(_csv(tmp_path, "A\n1\n2\n")).rows) == 2


def test_empty_formatted_cells_far_right_do_not_count(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["A", "B"])
    sheet.append([1, 2])
    sheet.cell(row=1, column=300).font = openpyxl.styles.Font(bold=True)   # empty, only formatted
    path = tmp_path / "formatiert.xlsx"
    book.save(path)
    assert load_table(path) == Table(["A", "B"], [["1", "2"]])
    assert load_table(_csv(tmp_path, "A;B;;;\n1;2;;;\n")) == Table(["A", "B"], [["1", "2"]])


def test_many_empty_lines_do_not_eat_memory(tmp_path):
    import resource
    before = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    table = load_table(_csv(tmp_path, "A\n1\n" + "\n" * 5_000_000))
    assert table.rows == [["1"]]
    assert resource.getrusage(resource.RUSAGE_SELF).ru_maxrss - before < 400_000   # KB


def test_chart_sheets_are_skipped(tmp_path):
    from openpyxl.chart import BarChart
    book = openpyxl.Workbook()
    data = book.active
    data.title = "Daten"
    data.append(["A"])
    data.append([1])
    chart = book.create_chartsheet("Diagramm", 0)
    chart.add_chart(BarChart())
    path = tmp_path / "diagramm.xlsx"
    book.save(path)
    assert sheet_names(path) == ["Daten"]
    assert load_table(path) == Table(["A"], [["1"]])


def test_hidden_first_sheet_is_not_the_default(tmp_path):
    book = openpyxl.Workbook()
    hidden = book.active
    hidden.title = "Versteckt"
    hidden.append(["H"])
    hidden.sheet_state = "hidden"
    shown = book.create_sheet("Sichtbar")
    shown.append(["S"])
    shown.append(["1"])
    path = tmp_path / "versteckt.xlsx"
    book.save(path)
    assert load_table(path) == Table(["S"], [["1"]])


@pytest.mark.parametrize("source", ["a\x00.csv", ""])
def test_bad_paths_give_series_errors(source):
    with pytest.raises(SeriesError):
        table_for(Series(source=source))


@pytest.mark.parametrize("value", ["a\x00.csv", ""])
def test_bad_paths_in_files_are_refused(tmp_path, value):
    import json
    data = storage.to_dict(Document("62"))
    data["document"]["series"] = {"source": value}
    path = tmp_path / "s.labelwerk"
    path.write_text(json.dumps(data))
    with pytest.raises(storage.StorageError, match="source"):
        storage.load(path)


def test_fill_shares_only_what_it_does_not_change():
    """Filled labels are read-only (rendering, printing): unchanged objects stay
    the same (renderer cache), changed ones are new – the original stays untouched."""
    doc = _doc(TextObject("{A}"), TextObject("fest"))
    filled = fill(doc, ["A"], ["1"])
    assert filled.objects[1] is doc.objects[1]
    assert filled.objects[0] is not doc.objects[0]
    assert filled.objects is not doc.objects
    assert doc.objects[0].text == "{A}"


def test_separator_is_chosen_over_the_first_lines(tmp_path):
    assert load_table(_csv(tmp_path, "Name, Vorname;Ort\nMeier, Hans;Berlin\n")).columns == [
        "Name, Vorname", "Ort"]
    assert load_table(_csv(tmp_path, '"Name; Firma",Ort\n"A; B",X\n')).columns == [
        "Name; Firma", "Ort"]
    assert load_table(_csv(tmp_path, "A;1\tB\n2;3\t4\n")).columns == ["A;1", "B"]


def test_utf16_text_from_excel(tmp_path):
    path = tmp_path / "unicode.txt"
    path.write_bytes("Klemme\tOrt\nX1\tKöln\n".encode("utf-16"))
    assert load_table(path) == Table(["Klemme", "Ort"], [["X1", "Köln"]])


def test_windows_characters_are_decoded(tmp_path):
    table = load_table(_csv(tmp_path, "Preis;Hinweis\n5 €;„gut“ – ja\n", "cp1252"))
    assert table.rows == [["5 €", "„gut“ – ja"]]


def test_too_long_cell_gives_a_german_message(tmp_path):
    with pytest.raises(SeriesError, match="Zelle zu lang"):
        load_table(_csv(tmp_path, "A\n" + "x" * 200_000 + "\n"))


def test_column_names_are_trimmed(tmp_path):
    assert load_table(_csv(tmp_path, " A ;B\n1;2\n")).columns == ["A", "B"]


def test_excel_number_formats_are_kept(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Nr", "Anteil", "Preis", "Zeit", "Groß"])
    sheet.append([7, 0.25, 12.5, datetime.time(9, 5), 1234567890123456])
    sheet["A2"].number_format = "000"
    sheet["B2"].number_format = "0%"
    sheet["C2"].number_format = "0.00"
    path = tmp_path / "formate.xlsx"
    book.save(path)
    assert load_table(path).rows == [["007", "25%", "12,50", "09:05", "1234567890123456"]]


def test_tabs_in_values_become_spaces():
    doc = _doc(TextObject("{A}"))
    assert fill(doc, ["A"], ["x\ty"]).objects[0].text == "x y"


def test_placeholders_may_have_spaces_inside_the_braces():
    doc = _doc(TextObject("{ Klemme }"))
    assert placeholders(doc) == ["Klemme"]
    assert fill(doc, ["Klemme"], ["X1"]).objects[0].text == "X1"


@pytest.mark.parametrize("numbering, message", [
    (Numbering(count=True), "Anzahl"), (Numbering(start=10**9, count=3), "zu groß"),
])
def test_numbering_bounds_also_for_results(numbering, message):
    with pytest.raises(SeriesError, match=message):
        numbering_table(numbering)


def test_check_refuses_wrong_types():
    with pytest.raises(SeriesError):
        series.check(Series(source=3))
    with pytest.raises(SeriesError):
        series.check(Series(numbering=Numbering(count=0)))


def test_a_fifo_is_not_read(tmp_path):
    fifo = tmp_path / "rohr.csv"
    os.mkfifo(fifo)
    with pytest.raises(SeriesError, match="keine Datei"):
        load_table(fifo)


def test_excel_xml_bombs_are_refused(tmp_path):
    path = _xlsx(tmp_path, {"Daten": [["A"], ["x"]]})
    bomb = ('<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">'
            '<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">]>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&lol2;</t></is></c></row>'
            '</sheetData></worksheet>')
    out = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = bomb.encode()
            dst.writestr(item, data)
    path.write_bytes(out.getvalue())
    with pytest.raises(SeriesError, match="nicht lesbar"):
        load_table(path)


def test_table_next_to_the_label_is_saved_relative(tmp_path):
    import json
    table = _csv(tmp_path, "A\n1\n")
    doc = Document("62", objects=[TextObject("{A}")], series=Series(source=str(table)))
    saved = storage.save(doc, tmp_path / "s.labelwerk")
    assert json.loads(saved.read_text())["document"]["series"]["source"] == "t.csv"
    assert storage.load(saved).series.source == str(table)
    elsewhere = tmp_path / "anders"
    elsewhere.mkdir()
    saved = storage.save(doc, elsewhere / "s.labelwerk")
    assert json.loads(saved.read_text())["document"]["series"]["source"] == str(table)


# --- Final review ----------------------------------------------------------------------------
def test_far_away_excel_rows_stop_early(tmp_path, monkeypatch):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["A1"], sheet["A2"], sheet["A200"] = "Klemme", "X1", "weit unten"
    path = tmp_path / "weit.xlsx"
    book.save(path)
    monkeypatch.setattr(series, "MAX_RAW_ROWS", 100)
    with pytest.raises(SeriesError, match="leere Zeilen"):
        load_table(path)


def test_table_knows_the_sheet_it_came_from(tmp_path):
    book = openpyxl.Workbook()
    hidden = book.active
    hidden.title = "Versteckt"
    hidden.sheet_state = "hidden"
    shown = book.create_sheet("Sichtbar")
    shown.append(["S"])
    shown.append([1])
    path = tmp_path / "t.xlsx"
    book.save(path)
    table = load_table(path)
    assert table.sheet == "Sichtbar" and table.sheets == ["Versteckt", "Sichtbar"]
    assert load_table(_csv(tmp_path, "A\n1\n")).sheet is None


def test_numbering_placeholder_in_every_language():
    from labelwerk import i18n
    from labelwerk.document import Numbering
    from labelwerk.series import fill, missing, numbering_table
    doc = Document("62", objects=[TextObject("{Nummer}/{Number}")])
    i18n.set_language("en")
    table = numbering_table(Numbering(count=2))
    assert table.columns == ["Number"] and missing(doc, table.columns) == []
    assert fill(doc, table.columns, table.rows[0]).objects[0].text == "1/1"
    i18n.set_language("de")
    table = numbering_table(Numbering(count=2))
    assert table.columns == ["Nummer"] and fill(doc, table.columns, ["7"]).objects[0].text == "7/7"
    assert missing(Document("62", objects=[TextObject("{Numero}")]), table.columns) == ["Numero"]


def test_unnamed_columns_in_every_language(tmp_path):
    from labelwerk import i18n
    from labelwerk.series import fill, load_table, missing
    path = tmp_path / "t.csv"
    path.write_text("Klemme;;Ort\nX1;24V;A\n", encoding="utf-8")
    i18n.set_language("en")
    table = load_table(path)
    assert table.columns == ["Klemme", "Column 2", "Ort"]
    doc = Document("62", objects=[TextObject("{Spalte 2} {Column 2}")])
    assert missing(doc, table.columns) == []
    assert fill(doc, table.columns, table.rows[0]).objects[0].text == "24V 24V"


# --- only the beginning (preview in the file dialog) --------------------------------------------------
@pytest.mark.parametrize("count, more", [(3, False), (4, True), (50, True)])
def test_load_table_reads_only_the_first_rows(tmp_path, count, more):
    path = tmp_path / "t.csv"
    path.write_text("A;B\n" + "".join(f"{i};x\n\n" for i in range(count)))   # empty rows do not count
    table = load_table(path, limit=3)
    assert table.columns == ["A", "B"] and [r[0] for r in table.rows] == ["0", "1", "2"][:count]
    assert table.more is more


def test_load_table_limit_in_excel(tmp_path):
    path = tmp_path / "t.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["A"])
    for i in range(10):
        book.active.append([i])
    book.save(path)
    table = load_table(path, limit=2)
    assert table.rows == [["0"], ["1"]] and table.more is True
    assert load_table(path).more is False and len(load_table(path).rows) == 10
