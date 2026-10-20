"""Series: tables from CSV/Excel or numbering, and placeholders ``{Column}`` in the text and in
the content of barcodes.

Qt-free. Table files are foreign input: size, rows and columns are limited (rows are
streamed, empty ones discarded at once), every error becomes a translated ``SeriesError`` –
never a traceback.
"""
from __future__ import annotations

import copy
import csv
import datetime
import io
import itertools
import os
import re
import stat
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field, replace
from pathlib import Path

from .document import (CUTS, MAX_DIGITS, MAX_NUMBER, MAX_SERIES_ROWS, BarcodeObject, Document,
                       Numbering, Paragraph, Series, Span, TextObject)
from . import i18n
from .i18n import NP_, N_, _, pgettext
from .storage import LINE_BREAKS

MAX_BYTES = 20 * 1024 * 1024        # table file
MAX_UNPACKED = 200 * 1024 * 1024    # Excel: unpacked (against zip bombs)
MAX_ROWS = MAX_SERIES_ROWS          # rows without column names (empty ones do not count)
MAX_COLUMNS = 200
MAX_CELL = 100_000                  # characters per cell (CSV)
MAX_RAW_ROWS = 100_000              # Excel: rows including empty ones (stray cell far below)
CSV_SUFFIXES = (".csv", ".tsv", ".txt")
EXCEL_SUFFIXES = (".xlsx", ".xlsm")
# Translators: name of the numbering column, used as placeholder {Number} on labels
NUMBER_COLUMN = NP_("series", "Number")   # column of the numbering – valid in every language
# Translators: column without a name in the table header: "Column 3"
COLUMN_WORD = NP_("series", "Column")     # column without a name in the header row: “Column 3”
# on a tie the earlier one wins: Tab hardly ever appears in text, German Excel uses ";"
_SEPARATORS = ("\t", ";", ",", "|")
_SNIFF_LINES = 20
_BREAKS = re.compile("[" + re.escape(LINE_BREAKS) + "\t]+")   # become spaces
_PLACEHOLDER = re.compile(r"\{([^{}\n]+)\}")


class SeriesError(ValueError):
    """Table or series unusable (translated message, with file name)."""


@dataclass
class Table:
    """Column names and rows (texts, every row as long as ``columns``); Excel: the sheet read
    and all worksheets (not counted in comparisons)."""
    columns: list[str]
    rows: list[list[str]]
    sheet: str | None = field(default=None, compare=False)
    sheets: list[str] = field(default_factory=list, compare=False)
    more: bool = field(default=False, compare=False)   # only the beginning read (limit), more follow


# --- Reading tables ----------------------------------------------------------------------------
def load_table(path: str | Path, sheet: str | None = None, limit: int | None = None) -> Table:
    """CSV or Excel; first row = column names; empty rows are dropped (also in the middle).
    ``sheet``: Excel sheet (None = first visible worksheet); ``limit``: read only that many
    rows (preview), ``Table.more`` says whether more follow."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        return _read_excel(path, sheet, limit)
    if suffix in CSV_SUFFIXES:
        return _read_csv(path, limit)
    raise SeriesError(_("{name}: not a table format – CSV or Excel (.xlsx)").format(
        name=path.name or path))


def sheet_names(path: str | Path) -> list[str]:
    """Worksheets of an Excel file (without chart sheets); CSV: none."""
    path = Path(path)
    if path.suffix.lower() not in EXCEL_SUFFIXES:
        return []
    book = _open_excel(path)
    try:
        return [ws.title for ws in book.worksheets]
    finally:
        book.close()


def _read_bytes(path: Path) -> bytes:
    """Read a file completely – only a regular one (a FIFO would hang), at most ``MAX_BYTES``."""
    try:
        info = os.stat(path)
        if not stat.S_ISREG(info.st_mode):
            raise SeriesError(_("{name}: not a file").format(name=path.name))
        if info.st_size > MAX_BYTES:
            raise SeriesError(_("{name}: file too large (at most {size} MB)").format(
                name=path.name, size=MAX_BYTES // 2**20))
        return path.read_bytes()
    except SeriesError:
        raise
    except FileNotFoundError as e:
        raise SeriesError(_("Table not found: {path}").format(path=path)) from e
    except (OSError, ValueError) as e:         # ValueError: NUL character in the path
        raise SeriesError(_("{name} cannot be read").format(name=path.name or path)) from e


def _decode(data: bytes, path: Path) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):          # “Unicode text” from Excel
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError as e:
            raise SeriesError(_("{name}: table not readable (UTF-16 damaged)").format(
                name=path.name)) from e
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")                             # accepts every byte


def _separator(text: str) -> str:
    """Separator with which the first lines yield the same number of fields (more than one)."""
    sample = "\n".join(itertools.islice((line for line in text.splitlines() if line.strip()),
                                        _SNIFF_LINES))
    best, best_score = ";", (False, 0)
    for separator in _SEPARATORS:
        try:
            rows = list(csv.reader(io.StringIO(sample), delimiter=separator))
        except csv.Error:
            continue
        if not rows or len(rows[0]) < 2:
            continue
        score = (True, sum(len(row) == len(rows[0]) for row in rows))
        if score > best_score:
            best, best_score = separator, score
    return best


def _read_csv(path: Path, limit: int | None = None) -> Table:
    text = _decode(_read_bytes(path), path)
    if "\x00" in text:                                   # binary file with the wrong extension
        raise SeriesError(_("{name}: table not readable (not a text file)").format(name=path.name))
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_separator(text))
    field_limit = csv.field_size_limit()
    try:
        csv.field_size_limit(MAX_CELL)
        return _collect(reader, path, limit)
    except csv.Error as e:
        if "field larger" in str(e):
            raise SeriesError(_("{name}: cell too long (at most {limit} characters)").format(
                name=path.name, limit=i18n.thousands(MAX_CELL))) from e
        raise SeriesError(_("{name}: table not readable").format(name=path.name)) from e
    finally:
        csv.field_size_limit(field_limit)


def _open_excel(path: Path):
    """Workbook from the bytes read once (no second opening of the path)."""
    import openpyxl
    if not openpyxl.DEFUSEDXML:                # without defusedxml, XML attacks would be possible
        raise SeriesError(_("{name}: reading Excel files is disabled without protection "
                            "(defusedxml missing)").format(name=path.name))
    data = io.BytesIO(_read_bytes(path))
    try:
        with zipfile.ZipFile(data) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED:
                raise SeriesError(_("{name}: file too large when unpacked").format(name=path.name))
        return openpyxl.load_workbook(data, read_only=True, data_only=True)
    except SeriesError:
        raise
    except Exception as e:  # openpyxl/zipfile raise different errors depending on the damage
        raise SeriesError(_("{name}: Excel file not readable").format(name=path.name)) from e


def _read_excel(path: Path, sheet: str | None, limit: int | None = None) -> Table:
    book = _open_excel(path)
    try:
        sheets = book.worksheets                          # without chart sheets
        if sheet is not None:
            chosen = next((ws for ws in sheets if ws.title == sheet), None)
            if chosen is None:
                raise SeriesError(_("{name}: sheet “{sheet}” missing").format(name=path.name,
                                                                              sheet=sheet))
        else:
            chosen = next((ws for ws in sheets if ws.sheet_state == "visible"),
                          sheets[0] if sheets else None)
        if chosen is None:
            raise SeriesError(_("{name}: table is empty").format(name=path.name))
        # empty rows cost time too: openpyxl yields every gap (cell in row 1 048 576 → 35 s)
        width = MAX_COLUMNS + 1
        if chosen.max_column:                             # dimension of the file, if given
            width = max(1, min(width, chosen.max_column))
        rows = ([_cell(cell.value, getattr(cell, "number_format", "General")) for cell in row]
                for row in _limited(chosen.iter_rows(max_col=width), path))
        table = _collect(rows, path, limit)
        table.sheet, table.sheets = chosen.title, [ws.title for ws in sheets]
        return table
    except SeriesError:
        raise
    except Exception as e:
        raise SeriesError(_("{name}: Excel file not readable").format(name=path.name)) from e
    finally:
        book.close()


def _limited(rows, path: Path):
    """At most ``MAX_RAW_ROWS`` rows (empty ones too) – otherwise a message instead of minutes of reading."""
    for number, row in enumerate(rows, start=1):
        if number > MAX_RAW_ROWS:
            raise SeriesError(_("{name}: too many empty rows (data up to row {limit}) – delete "
                                "empty rows at the end of the sheet").format(
                name=path.name, limit=i18n.thousands(MAX_RAW_ROWS)))
        yield row


def _cell(value, number_format: str = "General") -> str:
    """Cell value as German Excel shows it: leading zeros (“000”), percent, fixed
    decimal places, comma; date DD.MM.YYYY (with time only if not 0:00)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "WAHR" if value else "FALSCH"
    if isinstance(value, (int, float)):
        return _number(value, number_format or "General")
    if isinstance(value, datetime.datetime):
        return value.strftime("%d.%m.%Y" if value.time() == datetime.time() else "%d.%m.%Y %H:%M")
    if isinstance(value, datetime.date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M")
    if isinstance(value, datetime.timedelta):
        minutes = round(value.total_seconds() / 60)
        return f"{minutes // 60}:{minutes % 60:02d}"
    return str(value)


def _number(value: int | float, number_format: str) -> str:
    fmt = number_format.split(";")[0]                     # only the part for positive numbers
    if "%" in fmt:
        value, suffix = value * 100, "%"
    else:
        suffix = ""
    match = re.search(r"(0+)(?:\.(0+))?", fmt.replace("#", "").replace(",", ""))
    if match and fmt not in ("General", "@"):
        decimals = len(match.group(2) or "")
        point = i18n.content_decimal()
        text = f"{abs(value):.{decimals}f}"
        whole, __, fraction = text.partition(".")
        negative = value < 0 and float(text) != 0
        text = whole.zfill(len(match.group(1))) + (point + fraction if fraction else "")
        return ("-" if negative else "") + text + suffix
    if isinstance(value, int) or (value.is_integer() and abs(value) < 1e21):
        return f"{value:.0f}" + suffix if isinstance(value, float) else str(value) + suffix
    return f"{value:.15g}".replace(".", i18n.content_decimal()) + suffix


def _collect(rows: Iterable[list], path: Path, limit: int | None = None) -> Table:
    """Collect rows: cells as text (line breaks → spaces), empty cells at the end of a row
    and empty rows dropped, column names unique, check limits at once; ``limit``: stop after
    that many rows (``more``: another one follows)."""
    header: list[str] | None = None
    data: list[list[str]] = []
    more = False
    for row in rows:
        cells = [_BREAKS.sub(" ", str(cell)) for cell in row]
        while cells and not cells[-1].strip():
            cells.pop()
        if not cells:
            continue
        if header is None:
            if len(cells) > MAX_COLUMNS:
                raise SeriesError(_("{name}: too many columns (at most {limit})").format(
                    name=path.name, limit=MAX_COLUMNS))
            header = cells
            continue
        if limit is not None and len(data) == limit:
            more = True
            break
        if len(data) == MAX_ROWS:
            raise SeriesError(_("{name}: too many rows (at most {limit})").format(
                name=path.name, limit=i18n.thousands(MAX_ROWS)))
        data.append(cells)
    if header is None:
        raise SeriesError(_("{name}: table is empty").format(name=path.name))
    columns: list[str] = []
    for number, name in enumerate(header, start=1):
        name = name.strip() or f"{pgettext('series', COLUMN_WORD)} {number}"
        unique, n = name, 2
        while unique in columns:
            unique, n = f"{name} ({n})", n + 1
        columns.append(unique)
    width = len(columns)
    return Table(columns, [(row + [""] * width)[:width] for row in data], more=more)


# --- Numbering and series --------------------------------------------------------------------
def _whole(value) -> bool:
    return isinstance(value, int) and not isinstance(value, bool)


def check_numbering(numbering: Numbering) -> None:
    """Limits of the numbering – even the last number stays at most ``MAX_NUMBER``."""
    limit = i18n.thousands(MAX_NUMBER)
    for value, message in ((numbering.start, N_("Numbering start: whole number up to {limit}")),
                           (numbering.step, N_("Numbering step: whole number up to {limit}"))):
        if not _whole(value) or abs(value) > MAX_NUMBER:
            raise SeriesError(_(message).format(limit=limit))
    if numbering.step == 0:
        raise SeriesError(_("Numbering step must not be 0"))
    if not _whole(numbering.count) or not 1 <= numbering.count <= MAX_SERIES_ROWS:
        raise SeriesError(_("Numbering count: 1 to {limit}").format(
            limit=i18n.thousands(MAX_SERIES_ROWS)))
    if not _whole(numbering.digits) or not 0 <= numbering.digits <= MAX_DIGITS:
        raise SeriesError(_("Digits of the numbering: 0 to {limit}").format(limit=MAX_DIGITS))
    if abs(numbering.start + (numbering.count - 1) * numbering.step) > MAX_NUMBER:
        raise SeriesError(_("Numbering values too large (at most {limit})").format(limit=limit))


def numbering_table(numbering: Numbering) -> Table:
    """Column “Number” (name in the language, ``number_column``): ``start``, ``start + step`` … –
    padded with zeros to ``digits`` digits."""
    check_numbering(numbering)

    def text(n: int) -> str:
        digits = str(abs(n)).zfill(numbering.digits)
        return f"-{digits}" if n < 0 else digits
    return Table([number_column()], [[text(numbering.start + i * numbering.step)]
                                   for i in range(numbering.count)])


def check(value: Series) -> None:
    """Check a series: exactly one of table and numbering, known cut."""
    if (value.source is None) == (value.numbering is None):
        raise SeriesError(_("A series needs exactly one: table or numbering"))
    if value.source is not None and (not isinstance(value.source, str) or not value.source
                                     or "\x00" in value.source):
        raise SeriesError(_("Invalid series table: {source!r}").format(source=value.source))
    if value.sheet is not None and not isinstance(value.sheet, str):
        raise SeriesError(_("Invalid series sheet: {sheet!r}").format(sheet=value.sheet))
    if value.cut not in CUTS:
        raise SeriesError(_("Unknown cut {cut!r} – allowed: {allowed}").format(
            cut=value.cut, allowed=", ".join(CUTS)))
    if value.numbering is not None:
        check_numbering(value.numbering)


def table_for(value: Series, base: str | Path | None = None) -> Table:
    """Rows of the series: table (path relative to ``base``) or numbering."""
    check(value)
    if value.numbering is not None:
        return numbering_table(value.numbering)
    path = Path(value.source)
    if not path.is_absolute() and base is not None:
        path = Path(base) / path
    return load_table(path, value.sheet)


# --- Placeholders -------------------------------------------------------------------------------
def placeholders(doc: Document) -> list[str]:
    """Names of the placeholders ``{Name}`` in text objects and barcodes (in order, each once)."""
    names: list[str] = []
    for obj in doc.objects:
        text = (obj.text if isinstance(obj, TextObject) else obj.data
                if isinstance(obj, BarcodeObject) else "")
        for name in _PLACEHOLDER.findall(text):
            name = name.strip()                  # { Klemme } matches the column “Klemme”
            if name not in names:
                names.append(name)
    return names


def number_column() -> str:
    """Name of the numbering column in the language (“Nummer” / “Number”)."""
    return pgettext("series", NUMBER_COLUMN)


def _aliases(columns: list[str]) -> dict[str, str]:
    """Names in other languages → column: the numbering works equally well under every name
    (“{Nummer}” = “{Number}”), likewise columns without a name (“Spalte 3” = “Column 3”) –
    labels stay valid when someone changes the language."""
    numbers = i18n.all_translations("series", NUMBER_COLUMN)
    words = i18n.all_translations("series", COLUMN_WORD)
    out: dict[str, str] = {}
    for column in columns:
        if column in numbers:
            out.update({name: column for name in numbers if name not in columns})
        m = re.fullmatch(r"(\S+) (\d+)", column)
        if m and m[1] in words:
            out.update({f"{word} {m[2]}": column for word in words
                        if f"{word} {m[2]}" not in columns})
    return out


def missing(doc: Document, columns: list[str]) -> list[str]:
    """Placeholders without a matching column (numbering names in every language count)."""
    known = set(columns) | set(_aliases(columns))
    return [name for name in placeholders(doc) if name not in known]


def fill(doc: Document, columns: list[str], row: list[str], keep: int | None = None) -> Document:
    """``doc`` with the values of the row instead of the placeholders – **read-only use**
    (rendering, printing): objects without placeholders, images and ``keep`` are the same as in
    the original (renderer cache hits), only filled text and barcode objects are new; ``doc``
    stays unchanged.

    The value takes the formatting of the span in which its placeholder begins (also across
    span boundaries); unknown placeholders are left in place; line breaks and tabs in the
    value become spaces. ``keep``: object that stays unchanged (is being edited right now).
    """
    values = {name: _BREAKS.sub(" ", value) for name, value in zip(columns, row)}
    values.update({alias: values[column] for alias, column in _aliases(columns).items()
                   if column in values})
    out = copy.copy(doc)
    out.objects = [obj if i == keep else _fill_text(obj, values) if isinstance(obj, TextObject)
                   else _fill_barcode(obj, values) if isinstance(obj, BarcodeObject) else obj
                   for i, obj in enumerate(doc.objects)]
    return out


def _fill_barcode(obj: BarcodeObject, values: dict[str, str]) -> BarcodeObject:
    data = _PLACEHOLDER.sub(lambda m: values.get(m.group(1).strip(), m.group(0)), obj.data)
    return obj if data == obj.data else replace(obj, data=data)


def _fill_text(obj: TextObject, values: dict[str, str]) -> TextObject:
    paragraphs = [_fill_paragraph(p, values) for p in obj.paragraphs]
    if all(new is old for new, old in zip(paragraphs, obj.paragraphs)):
        return obj
    filled = copy.copy(obj)
    filled.paragraphs = paragraphs
    return filled


def _fill_paragraph(para: Paragraph, values: dict[str, str]) -> Paragraph:
    # characters with the index of their span; replace from the back so positions stay valid
    chars = [(c, i) for i, span in enumerate(para.spans) for c in span.text]
    text = "".join(c for c, __ in chars)
    matches = [m for m in _PLACEHOLDER.finditer(text) if m.group(1).strip() in values]
    if not matches:
        return para
    for m in reversed(matches):
        style = chars[m.start()][1]
        chars[m.start():m.end()] = [(c, style) for c in values[m.group(1).strip()]]
    spans: list[Span] = []
    styles: list[int] = []                          # span whose formatting each new one carries
    for c, i in chars:
        if styles and styles[-1] == i:
            spans[-1].text += c
        else:
            spans.append(replace(para.spans[i], text=c))
            styles.append(i)
    if not spans:                                   # only empty values: one empty line, formatted
        spans = [replace(para.spans[_span_at(para, matches[0].start())], text="")]
    return Paragraph(spans, para.align)


def _span_at(para: Paragraph, pos: int) -> int:
    """Index of the span in which position ``pos`` of the paragraph text lies."""

    for i, span in enumerate(para.spans):
        if pos < len(span.text):
            return i
        pos -= len(span.text)
    return len(para.spans) - 1
